import re, sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "docs", "VRChat改模完全教程.md")
OUT = os.path.join(ROOT, "docs", "VRChat改模完全教程.docx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)
GREY = RGBColor(0x44, 0x44, 0x44)


def set_run_font(run, name="微软雅黑", size=None, bold=None, color=None, mono=False):
    if mono:
        name = "Consolas"
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), name)
    rf.set(qn('w:hAnsi'), name)
    rf.set(qn('w:eastAsia'), '微软雅黑' if not mono else 'Consolas')
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_hyperlink(paragraph, url, text, size=10.5):
    part = paragraph.part
    r_id = part.relate_to(url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '2E74B5'); rPr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); rPr.append(u)
    rf = OxmlElement('w:rFonts')
    rf.set(qn('w:ascii'), '微软雅黑'); rf.set(qn('w:hAnsi'), '微软雅黑'); rf.set(qn('w:eastAsia'), '微软雅黑')
    rPr.append(rf)
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size*2))); rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve')
    new_run.append(t)
    hl.append(new_run)
    paragraph._p.append(hl)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')

def add_inline(paragraph, text, size=10.5, color=None):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            r = paragraph.add_run(tok[2:-2]); set_run_font(r, size=size, bold=True, color=color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = paragraph.add_run(tok[1:-1]); set_run_font(r, size=size-0.5, color=RGBColor(0xA3,0x15,0x15), mono=True)
        elif tok.startswith('[') and '](' in tok:
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            if m:
                add_hyperlink(paragraph, m.group(2), m.group(1), size=size)
            else:
                r = paragraph.add_run(tok); set_run_font(r, size=size, color=color)
        else:
            r = paragraph.add_run(tok); set_run_font(r, size=size, color=color)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=9.5, white=False):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    # strip bold markers for header
    txt = text.replace('**', '')
    r = p.add_run(txt)
    set_run_font(r, size=size, bold=bold, color=RGBColor(0xFF,0xFF,0xFF) if white else RGBColor(0x22,0x22,0x22))


def split_row(line):
    line = line.strip()
    if line.startswith('|'): line = line[1:]
    if line.endswith('|'): line = line[:-1]
    return [c.strip() for c in line.split('|')]


def main():
    with open(SRC, encoding='utf-8') as f:
        lines = f.read().split('\n')

    doc = Document()

    # base style
    st = doc.styles['Normal']
    st.font.name = '微软雅黑'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    st.paragraph_format.line_spacing = 1.32
    st.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)

    i = 0
    n = len(lines)
    first_h1_done = False

    while i < n:
        line = lines[i].rstrip()
        s = line.strip()

        # blank
        if not s:
            i += 1
            continue

        # code fence
        if s.startswith('```'):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            for bl in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(bl if bl else ' ')
                set_run_font(r, size=9, mono=True, color=RGBColor(0x1a,0x1a,0x1a))
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear')
                shd.set(qn('w:color'),'auto'); shd.set(qn('w:fill'),'F3F4F6')
                pPr.append(shd)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # horizontal rule
        if re.fullmatch(r'-{3,}', s):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'6')
            bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),'C9D3DF')
            pbdr.append(bottom); pPr.append(pbdr)
            i += 1
            continue

        # headings
        m = re.match(r'^(#{1,6})\s+(.*)$', s)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and not first_h1_done:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(180)
                p.paragraph_format.space_after = Pt(10)
                r = p.add_run(text.replace('**',''))
                set_run_font(r, size=30, bold=True, color=ACCENT)
                first_h1_done = True
            elif level == 1:
                doc.add_page_break()
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(8)
                r = p.add_run(text.replace('**',''))
                set_run_font(r, size=20, bold=True, color=ACCENT)
                pPr = p._p.get_or_add_pPr()
                pbdr = OxmlElement('w:pBdr')
                bottom = OxmlElement('w:bottom')
                bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'12')
                bottom.set(qn('w:space'),'4'); bottom.set(qn('w:color'),'1F4E79')
                pbdr.append(bottom); pPr.append(pbdr)
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(text.replace('**',''))
                set_run_font(r, size=15, bold=True, color=ACCENT2)
            elif level == 3:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(text.replace('**',''))
                set_run_font(r, size=12.5, bold=True, color=RGBColor(0x1a,0x1a,0x1a))
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(text.replace('**',''))
                set_run_font(r, size=11, bold=True, color=GREY)
            i += 1
            continue

        # table
        if s.startswith('|') and i + 1 < n and re.match(r'^\|[\s:\-\|]+\|?$', lines[i+1].strip()):
            header = split_row(s)
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            tbl = doc.add_table(rows=1, cols=len(header))
            tbl.style = 'Table Grid'
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl.autofit = True
            for j, h in enumerate(header):
                c = tbl.rows[0].cells[j]
                set_cell_text(c, h, bold=True, size=9.5, white=True)
                shade(c, '1F4E79')
            for row in rows:
                cells = tbl.add_row().cells
                for j in range(len(header)):
                    val = row[j] if j < len(row) else ''
                    set_cell_text(cells[j], val, size=9)
            # zebra
            for ri, row in enumerate(tbl.rows[1:], start=1):
                if ri % 2 == 0:
                    for c in row.cells:
                        shade(c, 'EDF2F8')
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # blockquote
        if s.startswith('>'):
            buf = []
            while i < n and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, ' '.join(x for x in buf if x), size=10, color=RGBColor(0x33,0x44,0x55))
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18')
            left.set(qn('w:space'),'8'); left.set(qn('w:color'),'2E74B5')
            pbdr.append(left); pPr.append(pbdr)
            continue

        # list items
        m = re.match(r'^(\s*)([-*+])\s+(.*)$', line)
        if m:
            indent = len(m.group(1))
            text = m.group(3)
            lvl = 1 if indent < 2 else 2
            if text.startswith('[ ] '):
                text = '☐ ' + text[4:]
            elif text.startswith('[x] '):
                text = '☑ ' + text[4:]
            else:
                text = '• ' + text
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55 * lvl + 0.2)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, text, size=10.5)
            i += 1
            continue

        m = re.match(r'^(\s*)(\d+)\.\s+(.*)$', line)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.25
            add_inline(p, f"{m.group(2)}. " + m.group(3), size=10.5)
            i += 1
            continue

        # itallic whole line (footer note)
        if s.startswith('*') and s.endswith('*') and not s.startswith('**') and len(s) > 3:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(s.strip('*'))
            set_run_font(r, size=9, color=RGBColor(0x77,0x77,0x77))
            i += 1
            continue

        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.35
        add_inline(p, s, size=10.5)
        i += 1

    doc.save(OUT)
    print("saved", OUT, os.path.getsize(OUT))


main()
